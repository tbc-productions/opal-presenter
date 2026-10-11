"""Editable DrawingML effects and objects for the portable slide designer."""
import math
from copy import deepcopy

from lxml import etree
from pptx.oxml.ns import qn


def _child(parent, tag, **attributes):
    return etree.SubElement(parent, qn(tag), {key: str(value) for key, value in attributes.items()})


def _colour(parent, colour, opacity=100):
    node = _child(parent, 'a:srgbClr', val=colour.lstrip('#').upper())
    if opacity != 100:
        _child(node, 'a:alpha', val=round(opacity * 1000))
    return node


def _outline(parent, properties, unit, opacity):
    for node in list(parent):
        if node.tag == qn('a:ln'):
            parent.remove(node)
    line = _child(parent, 'a:ln', w=round(properties.get('outline_width', 2) * unit))
    if properties.get('outline'):
        _colour(_child(line, 'a:solidFill'), properties.get('outline_colour', '#000000'), opacity)
        _child(line, 'a:prstDash', val='solid')
        _child(line, 'a:round')
    else:
        _child(line, 'a:noFill')
    return line


def _shadow(parent, properties, unit, opacity):
    for node in list(parent):
        if node.tag in (qn('a:effectLst'), qn('a:effectDag')):
            parent.remove(node)
    effects = _child(parent, 'a:effectLst')
    if properties.get('shadow'):
        x, y = properties.get('shadow_x', 4), properties.get('shadow_y', 4)
        shadow = _child(effects, 'a:outerShdw', blurRad=round(properties.get('shadow_blur', 4) * unit),
                        dist=round(math.hypot(x, y) * unit),
                        dir=round(math.degrees(math.atan2(y, x)) % 360 * 60000),
                        algn='ctr', rotWithShape='1')
        _colour(shadow, properties.get('shadow_colour', '#000000'), opacity)
    return effects


def apply_object_properties(shape, properties, unit, *, text=False):
    """Keep opacity, outlines, shadows, rotation and flips editable in PowerPoint."""
    properties = dict(properties)
    if 'outline' not in properties and properties.get('border_width') and not text:
        properties.update(outline=True, outline_width=properties['border_width'],
                          outline_colour=properties.get('border_colour', '#ffffff'))
    shape.rotation = properties.get('rotation', 0)
    xfrm = shape._element.spPr.get_or_add_xfrm()
    for key, attr in (('flip_horizontal', 'flipH'), ('flip_vertical', 'flipV')):
        if properties.get(key) and not text:
            xfrm.set(attr, '1')
    opacity = properties.get('opacity', 100)
    if text:
        if properties.get('flip_horizontal') or properties.get('flip_vertical'):
            # PowerPoint mirrors editable text using 3-D rotation, rather than a shape's flip flags.
            # https://support.microsoft.com/en-us/word/reverse-or-mirror-text
            for parent in (shape._element.spPr, shape.text_frame._txBody.bodyPr):
                scene = _child(parent, 'a:scene3d')
                camera = _child(scene, 'a:camera', prst='orthographicFront')
                _child(camera, 'a:rot', lat=10800000 if properties.get('flip_vertical') else 0,
                       lon=10800000 if properties.get('flip_horizontal') else 0, rev=0)
                _child(scene, 'a:lightRig', rig='flat', dir='t')
        # Text effects belong to each run; applying them to the box outlines its rectangle instead.
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                props = run._r.get_or_add_rPr()
                if 'outline' in properties:
                    line = _outline(props, properties, unit, opacity)
                    props.remove(line)
                    props.insert(0, line)
                fill = props.find(qn('a:solidFill'))
                if fill is not None and len(fill) and opacity != 100:
                    colour = fill[0]
                    existing = colour.find(qn('a:alpha'))
                    previous = int(existing.get('val')) if existing is not None else 100000
                    if existing is not None:
                        colour.remove(existing)
                    _child(colour, 'a:alpha', val=round(previous * opacity / 100))
                if 'shadow' in properties:
                    effects = _shadow(props, properties, unit, opacity)
                    props.remove(effects)
                    # Effect properties precede font declarations in CT_TextCharacterProperties.
                    index = next((i for i, node in enumerate(props)
                                  if node.tag in (qn('a:latin'), qn('a:ea'), qn('a:cs'))), len(props))
                    props.insert(index, effects)
    else:
        _outline(shape._element.spPr, properties, unit, opacity)
        _shadow(shape._element.spPr, properties, unit, opacity)
        for colour in shape._element.spPr.xpath('./a:solidFill/a:srgbClr'):
            if opacity != 100:
                _child(colour, 'a:alpha', val=round(opacity * 1000))
        for blip in shape._element.xpath('./p:blipFill/a:blip'):
            if opacity != 100:
                _child(blip, 'a:alphaModFix', amt=round(opacity * 1000))


def text_properties(theme, part, design=None):
    """Resolve inherited theme effects into pixel units before applying saved overrides."""
    prefix = 'font_' + part + '_'
    properties = {
        'outline': getattr(theme, prefix + 'outline', False),
        'outline_width': getattr(theme, prefix + 'outline_size', 0) * 4 / 3,
        'outline_colour': getattr(theme, prefix + 'outline_color', '#000000'),
        'shadow': getattr(theme, prefix + 'shadow', False),
        'shadow_x': getattr(theme, prefix + 'shadow_size', 0) * 4 / 3,
        'shadow_y': getattr(theme, prefix + 'shadow_size', 0) * 4 / 3,
        'shadow_blur': getattr(theme, prefix + 'shadow_blur', 0) * 4 / 3,
        'shadow_colour': getattr(theme, prefix + 'shadow_color', '#000000')}
    if design:
        scale = getattr(theme, 'song_design_output_height', design.get('canvas_height', 1080)) / \
            design.get('canvas_height', 1080)
        overrides = dict(design.get('lyrics', {}))
        for key in ('outline_width', 'shadow_x', 'shadow_y', 'shadow_blur'):
            if key in overrides:
                overrides[key] *= scale
        properties.update(overrides)
    return properties


def element_theme(theme, element, width, height, scale):
    """Reuse the rich-text writer for editable designer text boxes."""
    result = deepcopy(theme)
    for key, dimension in (('x', width), ('width', width), ('y', height), ('height', height)):
        setattr(result, 'font_main_' + key, element[key] * dimension / 100)
    for key, value in (('name', element.get('font', 'Arial')), ('size', element.get('font_size', 40) * scale),
                       ('color', element.get('colour', '#ffffff')), ('bold', element.get('bold', False)),
                       ('italics', element.get('italic', False))):
        setattr(result, 'font_main_' + key, value)
    result.display_horizontal_align = ('left', 'right', 'center').index(element.get('align', 'center'))
    result.display_vertical_align = 0
    return result
