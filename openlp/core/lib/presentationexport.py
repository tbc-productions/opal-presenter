# -*- coding: utf-8 -*-
##########################################################################
# OpenLP - Open Source Lyrics Projection                                 #
# Copyright (c) 2008 OpenLP Developers                                   #
# This program is free software: you can redistribute it and/or modify   #
# it under the terms of the GNU General Public License as published by   #
# the Free Software Foundation, either version 3 of the License, or      #
# (at your option) any later version.                                    #
##########################################################################
"""Export rendered service pages as editable PPTX or ODP presentations."""
import html
import base64
import mimetypes
from dataclasses import dataclass, replace
from io import BytesIO
from os import PathLike
from pathlib import Path
from uuid import uuid4
from zipfile import BadZipFile

from bs4 import BeautifulSoup, NavigableString
from PySide6 import QtCore, QtGui, QtSvg

from openlp.core.common.enum import ServiceItemType
from openlp.core.common.i18n import translate
from openlp.core.common.registry import Registry
from openlp.core.display.screens import ScreenList
from openlp.core.lib.formattingtags import FormattingTags
from openlp.core.lib import ItemCapabilities


@dataclass
class PresentationPage:
    theme: object = None
    text: str = ''
    footer: str = ''
    html_text: str = ''
    html_footer: str = ''
    image: Path = None
    logo: bool = False
    title: str = ''
    section_index: int = None
    section_name: str = ''
    media: Path = None
    media_auto_start: bool = False
    media_start_time: int = 0
    presentation: object = None
    slide_index: int = None


@dataclass(frozen=True)
class TextStyle:
    family: str
    size: float
    color: str
    bold: bool = False
    italic: bool = False
    underline: bool = False
    baseline: str = ''
    transform: str = ''
    opacity: float = 1.0


def _base_text_style(theme, part):
    prefix = 'font_' + part + '_'
    return TextStyle(family=getattr(theme, prefix + 'name'),
                     size=float(getattr(theme, prefix + 'size')),
                     color=QtGui.QColor(getattr(theme, prefix + 'color')).name(),
                     bold=bool(getattr(theme, prefix + 'bold')),
                     italic=bool(getattr(theme, prefix + 'italics')))


def _css_size(value, parent_size):
    value = value.strip().lower()
    try:
        if value.endswith('%'):
            return parent_size * float(value[:-1]) / 100
        if value.endswith('rem') or value.endswith('em'):
            return parent_size * float(value[:-3] if value.endswith('rem') else value[:-2])
        if value.endswith('pt'):
            return float(value[:-2]) * 4 / 3
        if value.endswith('px'):
            return float(value[:-2])
        return float(value)
    except ValueError:
        return parent_size


def _styled_lines(markup, base):
    """Turn OpenLP's rendered HTML into lines of editable, styled text runs."""
    lines = [[]]

    def line_break():
        lines.append([])

    def append(value, current):
        if current.transform == 'uppercase':
            value = value.upper()
        elif current.transform == 'lowercase':
            value = value.lower()
        elif current.transform == 'capitalize':
            value = value.title()
        for index, fragment in enumerate(value.split('\n')):
            if index:
                line_break()
            if fragment:
                if lines[-1] and lines[-1][-1][1] == current:
                    previous, _ = lines[-1][-1]
                    lines[-1][-1] = (previous + fragment, current)
                else:
                    lines[-1].append((fragment, current))

    def visit(node, current):
        if isinstance(node, NavigableString):
            append(str(node), current)
            return
        name = node.name.lower()
        if name in ('script', 'style'):
            return
        if name == 'br':
            line_break()
            return
        if name in ('p', 'div') and lines[-1]:
            line_break()
        if name in ('b', 'strong'):
            current = replace(current, bold=True)
        elif name in ('i', 'em'):
            current = replace(current, italic=True)
        elif name == 'u':
            current = replace(current, underline=True)
        elif name in ('sup', 'sub'):
            current = replace(current, baseline='super' if name == 'sup' else 'sub')
        if name == 'font':
            if node.get('face'):
                current = replace(current, family=node['face'])
            if node.get('color') and QtGui.QColor(node['color']).isValid():
                current = replace(current, color=QtGui.QColor(node['color']).name())
        declarations = {}
        for declaration in node.get('style', '').split(';'):
            key, separator, value = declaration.partition(':')
            if separator:
                declarations[key.strip().lower()] = value.strip()
        if declarations.get('display') == 'none' or declarations.get('visibility') == 'hidden':
            return
        color = declarations.get('-webkit-text-fill-color', declarations.get('color'))
        if color and QtGui.QColor(color).isValid():
            current = replace(current, color=QtGui.QColor(color).name())
        if 'font-family' in declarations:
            current = replace(current, family=declarations['font-family'].split(',')[0].strip(' \"\''))
        if 'font-size' in declarations:
            current = replace(current, size=_css_size(declarations['font-size'], current.size))
        if 'font-weight' in declarations:
            weight = declarations['font-weight'].lower()
            current = replace(current, bold=weight in ('bold', 'bolder') or weight.isdigit() and int(weight) >= 600)
        if 'font-style' in declarations:
            current = replace(current, italic=declarations['font-style'].lower() in ('italic', 'oblique'))
        if 'text-decoration' in declarations:
            current = replace(current, underline='underline' in declarations['text-decoration'].lower())
        if 'text-transform' in declarations:
            current = replace(current, transform=declarations['text-transform'].lower())
        if 'opacity' in declarations:
            try:
                current = replace(current, opacity=current.opacity * max(0, min(1, float(declarations['opacity']))))
            except ValueError:
                pass
        for child in node.children:
            visit(child, current)
        if name in ('p', 'div') and lines[-1]:
            line_break()

    for child in BeautifulSoup(markup or '', 'html.parser').contents:
        visit(child, base)
    while len(lines) > 1 and not lines[-1]:
        lines.pop()
    return lines


def _run_color(style, theme):
    """Approximate CSS opacity against the theme's base colour."""
    color = QtGui.QColor(style.color)
    if style.opacity < 1:
        background = QtGui.QColor(getattr(theme, 'background_color', '#000000'))
        color.setRgb(round(color.red() * style.opacity + background.red() * (1 - style.opacity)),
                     round(color.green() * style.opacity + background.green() * (1 - style.opacity)),
                     round(color.blue() * style.opacity + background.blue() * (1 - style.opacity)))
    return color.name()


def _plain_text(value):
    return html.unescape(BeautifulSoup(value or '', 'html.parser').get_text('\n')).strip()


def _expand_footer_tags(value):
    """Footers are stored as HTML, but can also contain OpenLP formatting tags."""
    if '{' not in value:
        return value
    for tag in FormattingTags.get_html_tags():
        value = value.replace(tag['start tag'], tag['start html'])
        if tag['end tag']:
            value = value.replace(tag['end tag'], tag['end html'])
    return value


def _progress(callback, message, completed=0, total=0, title=''):
    if callback:
        message = translate('OpenLP.PresentationExport', message)
        callback(message + ('\n' + title if title else ''), completed, total)


def service_pages(items, insert_logo=False, *, native_pptx=False, progress=None):
    """Use the same renderer and theme resolution as the live service."""
    pages = []
    items = list(items)
    presentations = {}
    for section_index, item in enumerate(items):
        _progress(progress, 'Preparing slides', section_index, len(items), item.title)
        item_pages = []
        source = None
        plugin = getattr(item, 'name', '')
        if item.service_item_type == ServiceItemType.Command and item.slides and plugin in ('media', 'presentations'):
            if not item.is_capable(ItemCapabilities.CanStream):
                source = Path(item.get_frame_path())
        if source and plugin == 'media':
            if not source.is_file():
                raise FileNotFoundError(f'The media file is missing: {source}')
            item_pages.append(PresentationPage(media=source, title=item.title,
                                               media_auto_start=item.will_auto_start,
                                               media_start_time=getattr(item, 'start_time', 0)))
        elif source and native_pptx and source.suffix.lower() == '.pptx':
            from openlp.core.lib.pptxmerge import PptxSource

            if source not in presentations:
                presentations[source] = PptxSource(source)
            presentation = presentations[source]
            if not presentation.slides:
                raise ValueError(f'The presentation has no slides: {source}')
            item_pages = [PresentationPage(presentation=presentation, slide_index=index, title=item.title)
                          for index in range(len(presentation.slides))]
        elif item.service_item_type == ServiceItemType.Text:
            theme = item.get_theme_data()
            for rendered, displayed in zip(item.rendered_slides, item.display_slides):
                footer_html = _expand_footer_tags(rendered.get('footer') or '')
                page_theme = item.get_slide_theme(rendered['design_key']) if rendered.get('design_key') else theme
                item_pages.append(PresentationPage(theme=page_theme, text=html.unescape(displayed['text']),
                                                   footer=_plain_text(footer_html),
                                                   html_text=rendered.get('text') or '',
                                                   html_footer=footer_html,
                                                   title=item.title))
        elif item.service_item_type in (ServiceItemType.Image, ServiceItemType.Command):
            for frame in item.slides:
                image = frame.get('path') if item.service_item_type == ServiceItemType.Image else frame.get('image')
                # Media commands use a QIcon for the UI, not a static thumbnail file.
                image = Path(image) if isinstance(image, (str, PathLike)) and image else None
                if image and image.is_file():
                    item_pages.append(PresentationPage(image=image, title=item.title))
        if not item_pages:
            # Streams and unavailable presentation thumbnails have no static frame.
            item_pages.append(PresentationPage(theme=item.get_theme_data(), text=item.title, title=item.title))
        for page in item_pages:
            page.section_index = section_index
            page.section_name = item.title
        pages.extend(item_pages)
        if insert_logo:
            pages.append(PresentationPage(logo=True, section_index=section_index,
                                          section_name=item.title))
        _progress(progress, 'Preparing slides', section_index + 1, len(items), item.title)
    return pages


def _page_sections(pages):
    """Keep repeated service-item names as distinct, ordered sections."""
    sections = []
    for index, page in enumerate(pages):
        if page.section_index is None:
            continue
        if not sections or sections[-1][0] != page.section_index:
            sections.append((page.section_index, page.section_name, [index]))
        else:
            sections[-1][2].append(index)
    return sections


def _load_image(path):
    path = str(path)
    if path.startswith('data:image/png;base64,'):
        return QtGui.QImage.fromData(base64.b64decode(path.split(',', 1)[1]))
    if path.lower().endswith('.svg'):
        renderer = QtSvg.QSvgRenderer(path)
        if renderer.isValid():
            size = renderer.defaultSize()
            image = QtGui.QImage(size, QtGui.QImage.Format.Format_ARGB32_Premultiplied)
            image.fill(QtCore.Qt.GlobalColor.transparent)
            painter = QtGui.QPainter(image)
            renderer.render(painter)
            painter.end()
            return image
    return QtGui.QImage(path)


def _image_bytes(image):
    data = QtCore.QByteArray()
    buffer = QtCore.QBuffer(data)
    buffer.open(QtCore.QIODevice.OpenModeFlag.WriteOnly)
    image.save(buffer, 'PNG')
    buffer.close()
    return bytes(data)


def _fit_image(painter, image, width, height, fit='cover'):
    if image.isNull():
        return
    if fit == 'stretch':
        painter.drawImage(QtCore.QRectF(0, 0, width, height), image)
        return
    if fit == 'repeat':
        painter.fillRect(QtCore.QRectF(0, 0, width, height), QtGui.QBrush(image))
        return
    scale = (max if fit == 'cover' else min)(width / image.width(), height / image.height())
    target_width, target_height = image.width() * scale, image.height() * scale
    rect = QtCore.QRectF((width - target_width) / 2, (height - target_height) / 2,
                         target_width, target_height)
    painter.drawImage(rect, image)


def _media_poster(page, aspect_ratio):
    """Use a video frame at presentation resolution, with a smooth fallback for unreadable media."""
    from openlp.core.lib.mediaposter import video_frame

    mime_type = mimetypes.guess_type(str(page.media))[0] or ''
    frame = QtGui.QImage()
    if not mime_type.startswith('audio/'):
        frame = video_frame(page.media, position_ms=page.media_start_time or 1000)
    width = max(1920, frame.width())
    height = max(1, round(width / aspect_ratio))
    poster = QtGui.QImage(width, height, QtGui.QImage.Format.Format_RGB32)
    poster.fill(QtGui.QColor('#000000'))
    painter = QtGui.QPainter(poster)
    painter.setRenderHint(QtGui.QPainter.RenderHint.SmoothPixmapTransform)
    painter.setRenderHint(QtGui.QPainter.RenderHint.Antialiasing)
    if not frame.isNull():
        _fit_image(painter, frame, width, height, fit='contain')
    else:
        size = min(width, height) * 0.06
        painter.setPen(QtCore.Qt.PenStyle.NoPen)
        painter.setBrush(QtGui.QColor('#dddddd'))
        painter.drawPolygon(QtGui.QPolygonF([
            QtCore.QPointF(width / 2 - size / 2, height / 2 - size),
            QtCore.QPointF(width / 2 - size / 2, height / 2 + size),
            QtCore.QPointF(width / 2 + size, height / 2)]))
    painter.end()
    return _image_bytes(poster)


def _background(theme, width, height, logo=False, logo_settings=None, *, filtered=True, force_image=False):
    """Rasterise non-solid or filtered fills while keeping foreground text editable."""
    if logo:
        if logo_settings is None:
            settings = Registry().get('settings')
            logo_settings = {key: settings.value('core/logo ' + key) for key in ('background color', 'file', 'fit')}
        color = logo_settings['background color']
        filename = logo_settings['file']
        fit = logo_settings['fit']
        kind = 'image'
    else:
        color = theme.background_color
        filename = getattr(theme, 'background_filename', None)
        fit = getattr(theme, 'background_fit', 'cover')
        kind = theme.background_type
    overlay = (getattr(theme, 'song_design', {}) or {}).get('background', {}) if not logo else {}
    design = (getattr(theme, 'song_design', {}) or {}) if not logo else {}
    filters = design.get('background_filters', {}) if filtered else {}
    if (kind == 'solid' or (kind not in ('image', 'gradient') and not logo)) and \
            not overlay.get('image') and not filters and not force_image:
        return str(color), None
    image = QtGui.QImage(width, height, QtGui.QImage.Format.Format_ARGB32)
    image.fill(QtGui.QColor(color))
    painter = QtGui.QPainter(image)
    if kind == 'gradient':
        direction = getattr(theme, 'background_direction', 'vertical')
        if direction == 'horizontal':
            gradient = QtGui.QLinearGradient(0, 0, width, 0)
        elif direction == 'leftTop':
            gradient = QtGui.QLinearGradient(0, 0, width, height)
        elif direction == 'leftBottom':
            gradient = QtGui.QLinearGradient(0, height, width, 0)
        elif direction == 'circular':
            gradient = QtGui.QRadialGradient(width / 2, height / 2, max(width, height) / 2)
        else:
            gradient = QtGui.QLinearGradient(0, 0, 0, height)
        gradient.setColorAt(0, QtGui.QColor(theme.background_start_color))
        gradient.setColorAt(1, QtGui.QColor(theme.background_end_color))
        painter.fillRect(0, 0, width, height, gradient)
    elif filename:
        _fit_image(painter, _load_image(filename), width, height, fit)
    if overlay.get('image'):
        painter.setOpacity(overlay.get('image_opacity', 100) / 100)
        _fit_image(painter, _load_image(overlay['image']), width, height, overlay.get('fit', 'cover'))
    painter.end()
    if filters:
        from openlp.core.lib.imagefilters import apply_filters
        scale = getattr(theme, 'song_design_output_height', height) / design.get('canvas_height', height)
        image, _ = apply_filters(image, filters, scale)
    return str(color), _image_bytes(image)


def _design_image(element, width, height, scale, *, bake_effects=False):
    """Rasterise one contain-fit image; include filter overflow without moving its centre."""
    from openlp.core.lib.imagefilters import apply_filters, add_shadow

    image = QtGui.QImage(max(1, round(element['width'] / 100 * width)),
                         max(1, round(element['height'] / 100 * height)), QtGui.QImage.Format.Format_ARGB32)
    image.fill(QtCore.Qt.GlobalColor.transparent)
    painter = QtGui.QPainter(image)
    painter.setRenderHint(QtGui.QPainter.RenderHint.SmoothPixmapTransform)
    border = 0
    if bake_effects:
        border = (element.get('outline_width', 2) if element.get('outline') else
                  element.get('border_width', 0) if 'outline' not in element else 0) * scale
        border = min(border, image.width() / 2, image.height() / 2)
    painter.translate(border, border)
    _fit_image(painter, _load_image(element.get('src', '')), max(1, image.width() - 2 * border),
               max(1, image.height() - 2 * border), 'contain')
    painter.resetTransform()
    if border:
        painter.setPen(QtGui.QPen(QtGui.QColor(element.get('outline_colour', element.get('border_colour', '#000000'))),
                                  border))
        painter.drawRect(QtCore.QRectF(border / 2, border / 2, image.width() - border, image.height() - border))
    painter.end()
    image, padding = apply_filters(image, element.get('filters'), scale, expand=True)
    if bake_effects:
        image, shadow_padding = add_shadow(image, element, scale)
        padding += shadow_padding
    return image, padding


def _geometry(theme, part, width, height):
    """Return relative textbox coordinates from OpenLP's pixel theme geometry."""
    x = max(0, float(getattr(theme, 'font_' + part + '_x'))) / width
    y = max(0, float(getattr(theme, 'font_' + part + '_y'))) / height
    w = max(1, float(getattr(theme, 'font_' + part + '_width'))) / width
    h = max(1, float(getattr(theme, 'font_' + part + '_height'))) / height
    return x, y, min(w, 1 - x), min(h, 1 - y)


def _hex_color(value):
    color = QtGui.QColor(str(value))
    return color.name().lstrip('#') if color.isValid() else '000000'


def export_presentation(items, destination, insert_logo=False, *, progress=None):
    """Write a presentation for service items. The suffix selects the format."""
    destination = Path(destination)
    pages = service_pages(items, insert_logo, native_pptx=destination.suffix.lower() == '.pptx', progress=progress)
    if not pages:
        raise ValueError('There are no items to export')
    geometry = ScreenList().current.display_geometry
    width, height = geometry.width(), geometry.height()
    if width <= 0 or height <= 0:
        raise ValueError('The display geometry is not configured')
    if destination.suffix.lower() == '.pptx':
        writer = _write_pptx
    elif destination.suffix.lower() == '.odp':
        writer = _write_odp
    else:
        raise ValueError('Choose a .pptx or .odp file')
    # Resolve settings on the GUI thread before starting the writer.
    logo_settings = None
    if progress and any(page.logo for page in pages):
        settings = Registry().get('settings')
        logo_settings = {key: settings.value('core/logo ' + key) for key in ('background color', 'file', 'fit')}
    if progress:
        from openlp.core.ui.presentationexportdialog import write_with_progress

        write_with_progress(lambda report: writer(pages, destination, width, height,
                                                  progress=report, logo_settings=logo_settings), progress)
    else:
        writer(pages, destination, width, height, logo_settings=logo_settings)


def export_items_dialog(parent, items, suffix, default_path, insert_logo=False):
    """Choose an output path for one presentation format and report errors."""
    from PySide6 import QtWidgets

    from openlp.core.common.i18n import translate
    from openlp.core.widgets.dialogs import FileDialog
    from openlp.core.ui.presentationexportdialog import ExportProgressDialog

    items = list(items)
    if not items:
        QtWidgets.QMessageBox.information(parent, translate('OpenLP.PresentationExport', 'Export presentation'),
                                          translate('OpenLP.PresentationExport', 'Select an item to export.'))
        return
    if suffix == '.pptx':
        file_filter = translate('OpenLP.PresentationExport', 'PowerPoint Presentation (*.pptx)')
    elif suffix == '.odp':
        file_filter = translate('OpenLP.PresentationExport', 'OpenDocument Presentation (*.odp)')
    else:
        raise ValueError('Choose a .pptx or .odp file')
    path, _ = FileDialog.getSaveFileName(
        parent, translate('OpenLP.PresentationExport', 'Export presentation'),
        default_path, file_filter)
    if not path:
        return
    if path.suffix.lower() != suffix:
        path = path.with_suffix(suffix)
    dialog = ExportProgressDialog(parent)
    dialog.show()
    error_message = None
    try:
        export_presentation(items, path, insert_logo, progress=dialog.update_progress)
    except (OSError, ValueError, RuntimeError, ImportError, BadZipFile) as error:
        error_message = str(error)
    finally:
        dialog.finish()
    if error_message is not None:
        QtWidgets.QMessageBox.critical(parent, translate('OpenLP.PresentationExport', 'Export failed'), error_message)
    elif parent is not None and hasattr(parent, 'statusBar'):
        parent.statusBar().showMessage(
            translate('OpenLP.PresentationExport', 'Presentation exported to {path}').format(path=path), 10000)


def _write_pptx(pages, destination, width, height, *, progress=None, logo_settings=None):
    from lxml import etree
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from openlp.core.lib.slidedesign import layer_order
    from openlp.core.lib.pptxdesign import apply_object_properties, text_properties, element_theme
    from openlp.core.lib.pptxfilters import native_filter_effects, apply_native_filters
    from pptx.oxml.ns import qn
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.util import Inches, Pt

    presentation = Presentation()
    presentation.core_properties.title = destination.stem
    presentation.core_properties.author = 'OpenLP+'
    presentation.core_properties.last_modified_by = 'OpenLP+'
    presentation.core_properties.comments = 'Exported from OpenLP+'
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = round(presentation.slide_width * height / width)
    source = next((page.presentation for page in pages if page.presentation is not None), None)
    if source is not None:
        # Keep original slide geometry exact for the first source deck. Generated pages adapt to it.
        presentation.slide_width, presentation.slide_height = source.size
    blank = presentation.slide_layouts[6]
    backgrounds = {}
    background_originals = {}
    originals = {}
    posters = {}

    def add_text(slide, markup, theme, part, properties=None):
        if not markup:
            return
        x, y, w, h = _geometry(theme, part, width, height)
        box = slide.shapes.add_textbox(round(presentation.slide_width * x),
                                       round(presentation.slide_height * y),
                                       round(presentation.slide_width * w),
                                       round(presentation.slide_height * h))
        frame = box.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        vertical = getattr(theme, 'display_vertical_align' + ('_footer' if part == 'footer' else ''), 0)
        frame.vertical_anchor = (MSO_ANCHOR.TOP, MSO_ANCHOR.MIDDLE, MSO_ANCHOR.BOTTOM)[vertical]
        align = getattr(theme, 'display_horizontal_align' + ('_footer' if part == 'footer' else ''), 0)
        base = _base_text_style(theme, part)
        if properties is not None:
            base = replace(base, size=base.size * 4 / 3)
        for index, line in enumerate(_styled_lines(markup, base)):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.alignment = (PP_ALIGN.LEFT, PP_ALIGN.RIGHT, PP_ALIGN.CENTER, PP_ALIGN.JUSTIFY)[align]
            paragraph.space_before = paragraph.space_after = Pt(0)
            for value, run_style in line:
                run = paragraph.add_run()
                run.text = value
                font = run.font
                font.name = run_style.family
                font.size = Pt(run_style.size * (presentation.slide_width / 914400) * 72 / width)
                font.bold = run_style.bold
                font.italic = run_style.italic
                font.underline = run_style.underline
                font.color.rgb = RGBColor.from_string(_hex_color(
                    run_style.color if properties is not None else _run_color(run_style, theme)).upper())
                if properties is not None and run_style.opacity < 1:
                    colour = run._r.get_or_add_rPr().find(qn('a:solidFill'))[0]
                    etree.SubElement(colour, qn('a:alpha'), val=str(round(run_style.opacity * 100000)))
                if run_style.baseline:
                    baseline = '30000' if run_style.baseline == 'super' else '-25000'
                    run._r.get_or_add_rPr().set('baseline', baseline)
        if properties is not None:
            apply_object_properties(box, properties, presentation.slide_width / width, text=True)
        return box

    def add_element(slide, theme, element, scale):
        original, effects = None, []
        properties = dict(element)
        for key in ('outline_width', 'border_width', 'shadow_x', 'shadow_y', 'shadow_blur'):
            if key in properties:
                properties[key] *= scale
        if element['kind'] == 'text':
            element_style = element_theme(theme, element, width, height, scale)
            box = add_text(slide, html.escape(element.get('text', '')).replace('\n', '<br>'),
                           element_style, 'main', properties)
            if box is not None:
                box.name = 'Text box'
            return
        x, y = round(element['x'] / 100 * presentation.slide_width), \
            round(element['y'] / 100 * presentation.slide_height)
        w, h = round(element['width'] / 100 * presentation.slide_width), \
            round(element['height'] / 100 * presentation.slide_height)
        if element['kind'] == 'image':
            filtered = bool(element.get('filters'))
            # CSS filters also affect the border; Office filters only affect image pixels.
            # Preserve the complete appearance when a filtered border cannot stay native.
            if not element.get('outline') and not element.get('border_width'):
                effects = native_filter_effects(element.get('filters'), scale)
            image, padding = _design_image(element, width, height, scale, bake_effects=filtered and not effects)
            if effects:
                from PIL import ImageOps
                from openlp.core.lib.imagefilters import to_pillow, to_qimage

                original, _ = _design_image(dict(element, filters={}), width, height, scale)
                if padding:
                    original = to_qimage(ImageOps.expand(to_pillow(original), padding, fill=(0, 0, 0, 0)))
            elif filtered:
                properties.update(outline=False, shadow=False, border_width=0)
            pad_x = round(padding * presentation.slide_width / width)
            pad_y = round(padding * presentation.slide_height / height)
            x, y, w, h = x - pad_x, y - pad_y, w + 2 * pad_x, h + 2 * pad_y
            shape = slide.shapes.add_picture(BytesIO(_image_bytes(image)), x, y, w, h)
        else:
            shape = slide.shapes.add_shape(MSO_SHAPE.OVAL if element['kind'] == 'ellipse' else MSO_SHAPE.RECTANGLE,
                                           x, y, w, h)
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string(element.get('fill', '#315efb').lstrip('#'))
        shape.name = element['kind'].capitalize()
        apply_object_properties(shape, properties, presentation.slide_width / width)
        if effects:
            apply_native_filters(shape, original, effects, originals)

    for page_index, page in enumerate(pages):
        _progress(progress, 'Embedding media' if page.media else 'Building presentation',
                  page_index, len(pages), page.title or page.section_name)
        slide = presentation.slides.add_slide(blank)
        if page.presentation is not None:
            # Replaced with the original slide and its full dependency graph after writing generated slides.
            continue
        if page.media:
            mime_type = mimetypes.guess_type(str(page.media))[0] or 'video/unknown'
            poster_key = (page.media, page.media_start_time)
            if poster_key not in posters:
                posters[poster_key] = _media_poster(page, presentation.slide_width / presentation.slide_height)
            movie = slide.shapes.add_movie(str(page.media), 0, 0, presentation.slide_width, presentation.slide_height,
                                           poster_frame_image=BytesIO(posters[poster_key]), mime_type=mime_type)
            if mime_type.startswith('audio/'):
                audio_file = movie._element.xpath('.//a:videoFile')[0]
                video_rid = audio_file.get(qn('r:link'))
                media_part = slide.part.related_part(video_rid)
                audio_file.tag = qn('a:audioFile')
                audio_file.set(qn('r:link'), slide.part.relate_to(media_part, RT.AUDIO))
                slide.part.drop_rel(video_rid)
                slide._element.xpath('.//p:video')[0].tag = qn('p:audio')
            if page.media_auto_start:
                for condition in slide._element.xpath('.//p:cMediaNode/p:cTn/p:stCondLst/p:cond'):
                    condition.set('delay', '0')
            continue
        if page.image:
            slide.shapes.add_picture(BytesIO(_image_bytes(_load_image(page.image))), 0, 0,
                                     presentation.slide_width, presentation.slide_height)
            continue
        key = 'logo' if page.logo else id(page.theme)
        if key not in backgrounds:
            backgrounds[key] = _background(page.theme, width, height, page.logo, logo_settings)
        color, image = backgrounds[key]
        if image:
            background = slide.shapes.add_picture(BytesIO(image), 0, 0,
                                                  presentation.slide_width, presentation.slide_height)
            design = getattr(page.theme, 'song_design', {}) or {} if not page.logo else {}
            scale = getattr(page.theme, 'song_design_output_height', height) / design.get('canvas_height', height)
            effects = native_filter_effects(design.get('background_filters'), scale)
            if effects:
                if key not in background_originals:
                    _, original = _background(page.theme, width, height, filtered=False, force_image=True)
                    background_originals[key] = QtGui.QImage.fromData(original)
                apply_native_filters(background, background_originals[key], effects, originals)
        else:
            fill = slide.background.fill
            fill.solid()
            fill.fore_color.rgb = RGBColor.from_string(_hex_color(color).upper())
        if page.theme:
            design = getattr(page.theme, 'song_design', {}) or {}
            elements = {element['id']: element for element in design.get('elements', [])}
            scale = getattr(page.theme, 'song_design_output_height', height) / design.get('canvas_height', height)
            for identifier in layer_order(design):
                if identifier == 'lyrics':
                    box = add_text(slide, page.html_text or html.escape(page.text).replace('\n', '<br>'),
                                   page.theme, 'main', text_properties(page.theme, 'main', design) if design else None)
                    if box is not None:
                        box.name = 'Slide text'
                else:
                    add_element(slide, page.theme, elements[identifier], scale)
            add_text(slide, page.html_footer or html.escape(page.footer).replace('\n', '<br>'), page.theme, 'footer',
                     text_properties(page.theme, 'footer') if design else None)
    _progress(progress, 'Building presentation', len(pages), len(pages))
    sections = _page_sections(pages)
    if sections:
        presentation_ns = 'http://schemas.openxmlformats.org/presentationml/2006/main'
        section_ns = 'http://schemas.microsoft.com/office/powerpoint/2010/main'
        root = presentation._element
        extension_list = root.find(f'{{{presentation_ns}}}extLst')
        if extension_list is None:
            extension_list = etree.SubElement(root, f'{{{presentation_ns}}}extLst')
        extension = etree.SubElement(extension_list, f'{{{presentation_ns}}}ext',
                                     uri='{521415D9-36F7-43E2-AB2F-B90AF26B5E84}')
        section_list = etree.SubElement(extension, f'{{{section_ns}}}sectionLst',
                                        nsmap={'p14': section_ns})
        slide_ids = presentation.slides._sldIdLst
        for _, name, indices in sections:
            section = etree.SubElement(section_list, f'{{{section_ns}}}section',
                                       name=name, id='{' + str(uuid4()).upper() + '}')
            id_list = etree.SubElement(section, f'{{{section_ns}}}sldIdLst')
            for index in indices:
                etree.SubElement(id_list, f'{{{section_ns}}}sldId', id=str(slide_ids[index].id))
    if source is None:
        _progress(progress, 'Writing file', title=destination.name)
        presentation.save(str(destination))
    else:
        from openlp.core.lib.pptxmerge import PptxSlideImporter

        generated = BytesIO()
        _progress(progress, 'Packaging presentation')
        presentation.save(generated)
        importer = PptxSlideImporter(generated.getvalue())
        groups = {}
        for index, page in enumerate(pages):
            if page.presentation is not None:
                key = (page.presentation, page.section_index)
                groups.setdefault(key, []).append((page.slide_index, index))
        for occurrence, ((source, _), indices) in enumerate(groups.items(), 1):
            _progress(progress, 'Copying original slides', occurrence - 1, len(groups),
                      pages[indices[0][1]].title)
            try:
                importer.import_slides(source, indices, occurrence)
            except (KeyError, etree.XMLSyntaxError) as error:
                raise ValueError(f'Unable to copy the source PowerPoint slides: {error}') from error
        _progress(progress, 'Writing file', title=destination.name)
        importer.save(destination)


def _write_odp(pages, destination, width, height, *, progress=None, logo_settings=None):
    from odf import dc, draw, style, text
    from odf.element import Element
    from odf.namespaces import LOEXTNS
    from odf.opendocument import OpenDocumentPresentation
    from openlp.core.lib.imagefilters import to_pillow, to_qimage
    from openlp.core.lib.slidedesign import layer_order

    document = OpenDocumentPresentation()
    document.meta.addElement(dc.Title(text=destination.stem))
    document.meta.addElement(dc.Description(text='Exported from OpenLP+'))
    slide_width = 33.866  # 13.333 inches in cm
    slide_height = slide_width * height / width
    layout = style.PageLayout(name='OpenLPLayout')
    layout.addElement(style.PageLayoutProperties(pagewidth=f'{slide_width:.3f}cm',
                                                 pageheight=f'{slide_height:.3f}cm',
                                                 printorientation='landscape'))
    document.automaticstyles.addElement(layout)
    text_styles = {}
    run_styles = {}
    backgrounds = {}
    masters = {}
    page_names = []

    def add_text(page, markup, theme, part):
        if not markup:
            return
        key = (id(theme), part)
        if key not in text_styles:
            graphic = style.Style(name=f'OpenLPBox{len(text_styles)}', family='graphic')
            vertical = getattr(theme, 'display_vertical_align' + ('_footer' if part == 'footer' else ''), 0)
            graphic.addElement(style.GraphicProperties(
                fill='none', stroke='none',
                textareaverticalalign=('top', 'middle', 'bottom')[vertical]))
            document.automaticstyles.addElement(graphic)
            paragraph = style.Style(name=f'OpenLPText{len(text_styles)}', family='paragraph')
            align = getattr(theme, 'display_horizontal_align' + ('_footer' if part == 'footer' else ''), 0)
            paragraph.addElement(style.ParagraphProperties(
                textalign=('start', 'end', 'center', 'justify')[align], margin='0cm'))
            point_size = (float(getattr(theme, 'font_' + part + '_size')) *
                          slide_width / width * 72 / 2.54)
            paragraph.addElement(style.TextProperties(
                fontfamily=getattr(theme, 'font_' + part + '_name'),
                fontsize=f'{point_size:.2f}pt',
                color='#' + _hex_color(getattr(theme, 'font_' + part + '_color')),
                fontweight='bold' if getattr(theme, 'font_' + part + '_bold') else 'normal',
                fontstyle='italic' if getattr(theme, 'font_' + part + '_italics') else 'normal'))
            document.automaticstyles.addElement(paragraph)
            text_styles[key] = graphic, paragraph
        graphic, paragraph = text_styles[key]
        x, y, w, h = _geometry(theme, part, width, height)
        frame = draw.Frame(stylename=graphic, x=f'{x * slide_width:.3f}cm',
                           y=f'{y * slide_height:.3f}cm', width=f'{w * slide_width:.3f}cm',
                           height=f'{h * slide_height:.3f}cm')
        box = draw.TextBox()
        for line in _styled_lines(markup, _base_text_style(theme, part)):
            paragraph_element = text.P(stylename=paragraph)
            for value, run_style in line:
                run_key = (run_style, _run_color(run_style, theme))
                if run_key not in run_styles:
                    run_style_element = style.Style(name=f'OpenLPRun{len(run_styles)}', family='text')
                    properties = dict(fontfamily=run_style.family,
                                      fontsize=f'{run_style.size * slide_width / width * 72 / 2.54:.2f}pt',
                                      color=run_key[1],
                                      fontweight='bold' if run_style.bold else 'normal',
                                      fontstyle='italic' if run_style.italic else 'normal')
                    if run_style.underline:
                        properties.update(textunderlinestyle='solid', textunderlinetype='single')
                    if run_style.baseline:
                        properties['textposition'] = ('super 58%' if run_style.baseline == 'super' else 'sub 58%')
                    run_style_element.addElement(style.TextProperties(**properties))
                    document.automaticstyles.addElement(run_style_element)
                    run_styles[run_key] = run_style_element
                paragraph_element.addElement(text.Span(stylename=run_styles[run_key], text=value))
            if not line:
                paragraph_element.addText(' ')
            box.addElement(paragraph_element)
        frame.addElement(box)
        page.addElement(frame)

    def add_image(page, element, scale):
        from PIL import Image, ImageOps

        image, _ = _design_image(element, width, height, scale, bake_effects=True)
        pixels = to_pillow(image)
        if element.get('flip_horizontal'):
            pixels = ImageOps.mirror(pixels)
        if element.get('flip_vertical'):
            pixels = ImageOps.flip(pixels)
        if element.get('rotation'):
            pixels = pixels.rotate(-element['rotation'], resample=Image.Resampling.BICUBIC, expand=True)
        opacity = element.get('opacity', 100) / 100
        if opacity != 1:
            pixels.putalpha(pixels.getchannel('A').point(lambda alpha: round(alpha * opacity)))
        # Expanded blur/shadow/rotation bounds stay centred on the original image box.
        x = (element['x'] + element['width'] / 2) / 100 * width - pixels.width / 2
        y = (element['y'] + element['height'] / 2) / 100 * height - pixels.height / 2
        picture = document.addPictureFromString(_image_bytes(to_qimage(pixels)), 'image/png')
        frame = draw.Frame(name='Image', x=f'{x / width * slide_width:.3f}cm',
                           y=f'{y / height * slide_height:.3f}cm',
                           width=f'{pixels.width / width * slide_width:.3f}cm',
                           height=f'{pixels.height / height * slide_height:.3f}cm')
        frame.addElement(draw.Image(href=picture))
        page.addElement(frame)

    for index, source in enumerate(pages):
        _progress(progress, 'Embedding media' if source.media else 'Building presentation',
                  index, len(pages), source.title or source.section_name)
        key = 'logo' if source.logo else ('image', index) if source.image or source.media else id(source.theme)
        if key not in backgrounds:
            backgrounds[key] = (_background(source.theme, width, height, source.logo, logo_settings)
                                if not source.image and not source.media else None)
        if key not in masters:
            fill_style = style.Style(name=f'OpenLPPage{index}', family='drawing-page')
            color = backgrounds[key][0] if backgrounds[key] else '#000000'
            fill_style.addElement(style.DrawingPageProperties(fill='solid', fillcolor=color))
            document.automaticstyles.addElement(fill_style)
            master = style.MasterPage(name=f'OpenLPMaster{index}', pagelayoutname=layout)
            document.masterstyles.addElement(master)
            masters[key] = master, fill_style
        master, fill_style = masters[key]
        page = draw.Page(name=f'Slide {index + 1}', masterpagename=master,
                         stylename=fill_style)
        page_names.append(page.getAttribute('name'))
        document.presentation.addElement(page)
        if source.media:
            mime_type = mimetypes.guess_type(str(source.media))[0] or 'video/unknown'
            media_path = f'Media/{index}{source.media.suffix}'
            document.addPicture(media_path, mediatype=mime_type, content=source.media.read_bytes())
            frame = draw.Frame(x='0cm', y='0cm', width=f'{slide_width:.3f}cm', height=f'{slide_height:.3f}cm')
            player = draw.Plugin(href=media_path, mimetype='application/vnd.sun.star.media')
            player.addElement(draw.Param(name='Loop', value='false'))
            frame.addElement(player)
            page.addElement(frame)
            continue
        image = None
        if source.image:
            image = _image_bytes(_load_image(source.image))
        elif backgrounds[key]:
            image = backgrounds[key][1]
        if image:
            picture = document.addPictureFromString(image, 'image/png')
            frame = draw.Frame(x='0cm', y='0cm', width=f'{slide_width:.3f}cm',
                               height=f'{slide_height:.3f}cm')
            frame.addElement(draw.Image(href=picture))
            page.addElement(frame)
        if source.theme:
            design = getattr(source.theme, 'song_design', {}) or {}
            elements = {element['id']: element for element in design.get('elements', [])}
            scale = getattr(source.theme, 'song_design_output_height', height) / design.get('canvas_height', height)
            for identifier in layer_order(design):
                if identifier == 'lyrics':
                    add_text(page, source.html_text or html.escape(source.text).replace('\n', '<br>'),
                             source.theme, 'main')
                elif elements[identifier]['kind'] == 'image':
                    add_image(page, elements[identifier], scale)
            add_text(page, source.html_footer or html.escape(source.footer).replace('\n', '<br>'),
                     source.theme, 'footer')
    _progress(progress, 'Building presentation', len(pages), len(pages))
    sections = _page_sections(pages)
    if sections:
        section_list = Element(qname=(LOEXTNS, 'section-list'), check_grammar=False)
        for _, name, indices in sections:
            section = Element(qname=(LOEXTNS, 'section'), check_grammar=False)
            section.setAttrNS(LOEXTNS, 'name', name)
            section.setAttrNS(LOEXTNS, 'identifier', '{' + str(uuid4()).upper() + '}')
            for index in indices:
                section_slide = Element(qname=(LOEXTNS, 'section-slide'), check_grammar=False)
                section_slide.setAttrNS(LOEXTNS, 'name', page_names[index])
                section.addElement(section_slide, check_grammar=False)
            section_list.addElement(section, check_grammar=False)
        document.presentation.addElement(section_list, check_grammar=False)
    _progress(progress, 'Writing file', title=destination.name)
    document.save(str(destination))
