"""Validate editable Office effects, embedded originals and raster fallbacks."""
import base64
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile

import numpy
import pytest
from imagecodecs import jpegxr_decode
from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn
from PySide6 import QtGui

from opal.core.lib.pptxfilters import native_filter_effects, NAMESPACE, PHOTO_REL, BGRA_FORMAT, _office_jpegxr
from opal.core.lib.presentationexport import PresentationPage, _image_bytes, _write_pptx
from tests.openlp_core.lib.test_presentationexport import _theme


@pytest.mark.parametrize('filters', [
    {}, {'brightness': 100}, {'unknown': 10}, {'invert': 20}, {'sepia': 80}, {'hue-rotate': 45},
    {'brightness': 50, 'invert': 10}, {'blur': 60},
])
def test_unsupported_or_neutral_stacks_do_not_get_native_controls(filters):
    assert native_filter_effects(filters, scale=2) == []


def test_native_controls_preserve_order_ranges_and_combine_saturation():
    assert native_filter_effects({'blur': 5, 'brightness': 125, 'contrast': 80,
                                  'grayscale': 25, 'saturate': 120}, scale=.5) == [
        ('artisticBlur', {'radius': 2}),
        ('brightnessContrast', {'bright': 25000, 'contrast': -20000}),
        ('saturation', {'sat': 90000})]
    assert native_filter_effects({'grayscale': 100}) == [('saturation', {'sat': 0})]


def _picture():
    image = QtGui.QImage(80, 40, QtGui.QImage.Format.Format_ARGB32)
    image.fill(QtGui.QColor('#80ff0000'))
    return 'data:image/png;base64,' + base64.b64encode(_image_bytes(image)).decode()


def _original(shape):
    layer = shape._element.find('.//{' + NAMESPACE + '}imgLayer')
    relationship = shape.part.rels[layer.get(qn('r:embed'))]
    assert relationship.reltype == PHOTO_REL
    assert relationship.target_part.content_type == 'image/vnd.ms-photo'
    assert BGRA_FORMAT in relationship.target_part.blob
    return layer, jpegxr_decode(relationship.target_part.blob)


@pytest.mark.parametrize('alpha', [0, 127, 255, None])
def test_office_original_preserves_sharp_colour_and_alpha_pixels(alpha):
    # Noise and alternating edges catch channel swaps, premultiplication and loss
    # of detail; constant-colour sources cannot detect a baked blur regression.
    pixels = numpy.random.default_rng(37).integers(0, 256, (32, 48, 4), dtype=numpy.uint8)
    if alpha is not None:
        pixels[:, :, 3] = alpha
    blob = _office_jpegxr(Image.fromarray(pixels))

    assert BGRA_FORMAT in blob  # PowerPoint's supported decoded pixel format.
    numpy.testing.assert_array_equal(jpegxr_decode(blob), pixels)


def test_native_images_keep_editable_effects_lossless_originals_and_cached_pixels(tmp_path):
    theme = _theme()
    theme.song_design = {'canvas_height': 180, 'elements': [
        {'id': 'picture', 'kind': 'image', 'src': _picture(), 'x': 20, 'y': 20, 'width': 25, 'height': 25,
         'filters': {'brightness': 50, 'saturate': 0}, 'opacity': 40, 'rotation': 30,
         'shadow': True, 'flip_horizontal': True}]}
    original_design = deepcopy(theme.song_design)
    target = tmp_path / 'native.pptx'
    _write_pptx([PresentationPage(theme=theme, text='Editable text')] * 2, target, 320, 180)

    deck = Presentation(str(target))
    shape = next(shape for shape in deck.slides[0].shapes if shape.name == 'Image')
    layer, original = _original(shape)
    effects = [(child[0].tag.split('}')[1], dict(child[0].attrib)) for child in layer]
    assert effects == [('brightnessContrast', {'bright': '-50000', 'contrast': '0'}),
                       ('saturation', {'sat': '0'})]
    assert tuple(original[22, 40]) == (255, 0, 0, 128)
    cached = Image.open(BytesIO(shape.image.blob)).convert('RGBA')
    assert cached.getpixel((40, 22)) == (27, 27, 27, 128)
    assert shape.rotation == 30
    assert shape._element.spPr.xfrm.get('flipH') == '1'
    assert shape._element.xpath('./p:spPr/a:effectLst/a:outerShdw')
    blip = shape._element.blipFill.blip
    assert blip[-1].tag == qn('a:extLst')
    assert blip.find(qn('a:alphaModFix')).get('amt') == '40000'
    assert not blip.findall(qn('a:lum'))  # The cached PNG already contains the corrections.
    with ZipFile(target) as package:
        assert len([name for name in package.namelist() if name.endswith('.wdp')]) == 1
    assert theme.song_design == original_design
    # Check the extra relationship survives the package reader/writer round trip.
    deck.save(str(tmp_path / 'roundtrip.pptx'))
    reopened_shape = next(shape for shape in Presentation(str(tmp_path / 'roundtrip.pptx')).slides[1].shapes
                          if shape.name == 'Image')
    _, reopened = _original(reopened_shape)
    numpy.testing.assert_array_equal(original, reopened)


def test_native_background_blur_keeps_original_and_exact_cached_bounds(tmp_path):
    theme = _theme()
    theme.song_design = {'canvas_height': 180, 'background_filters': {'blur': 4, 'grayscale': 100}}
    target = tmp_path / 'background.pptx'
    _write_pptx([PresentationPage(theme=theme, text='Editable lyrics')], target, 320, 180)
    deck = Presentation(str(target))
    background, lyrics = deck.slides[0].shapes
    layer, original = _original(background)
    assert layer[0][0].get('radius') == '4'
    assert layer[1][0].get('sat') == '0'
    assert tuple(original[90, 160]) == (18, 52, 86, 255)
    assert original.shape[:2] == (180, 320)
    assert background.width == deck.slide_width
    assert background.height == deck.slide_height
    assert lyrics.text == 'Editable lyrics'


def test_native_image_blur_padding_preserves_position_and_original_size(tmp_path):
    theme = _theme()
    theme.song_design = {'canvas_height': 180, 'elements': [
        {'id': 'picture', 'kind': 'image', 'src': _picture(), 'x': 20, 'y': 20, 'width': 25, 'height': 25,
         'filters': {'blur': 4}}]}
    target = tmp_path / 'blur.pptx'
    _write_pptx([PresentationPage(theme=theme)], target, 320, 180)
    deck = Presentation(str(target))
    shape = deck.slides[0].shapes[0]
    _, original = _original(shape)
    cached = Image.open(BytesIO(shape.image.blob))
    assert cached.size == (104, 69)  # 80 x 45 plus 12 pixels on each edge.
    assert original.shape[:2] == (69, 104)
    assert original[0, 0, 3] == 0
    assert (shape.left + shape.width / 2) / deck.slide_width == pytest.approx(.325, abs=1e-6)


@pytest.mark.parametrize('properties', [
    {'filters': {'brightness': 50, 'sepia': 100}},
    {'filters': {'blur': 2, 'invert': 100}},
    {'filters': {'brightness': 50}, 'outline': True, 'outline_width': 2},
])
def test_unsupported_combinations_bake_once_without_broken_editable_metadata(tmp_path, properties):
    theme = _theme()
    theme.song_design = {'elements': [
        dict(id='picture', kind='image', src=_picture(), x=20, y=20, width=25, height=25, **properties)]}
    target = tmp_path / 'fallback.pptx'
    _write_pptx([PresentationPage(theme=theme)], target, 320, 180)
    with ZipFile(target) as package:
        assert not any(name.endswith('.wdp') for name in package.namelist())
    shape = Presentation(str(target)).slides[0].shapes[0]
    assert shape._element.find('.//{' + NAMESPACE + '}imgLayer') is None
    assert Image.open(BytesIO(shape.image.blob)).getpixel((40, 22)) != (255, 0, 0, 128)
