"""Package-level checks for editable presentation export."""
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from PySide6 import QtGui

from openlp.core.common.enum import ServiceItemType
from openlp.core.display import render as render_module
from openlp.core.lib.formattingtags import FormattingTags
from openlp.core.lib.presentationexport import PresentationPage, _styled_lines, _write_odp, _write_pptx, \
    service_pages, _base_text_style
from openlp.core.ui.servicemanager import ServiceManager
from openlp.core.widgets.dialogs import FileDialog
from tests.openlp_core.ui.test_mainwindow import main_window_reduced as _main_window_reduced
from tests.utils.constants import RESOURCE_PATH


presentation_main_window = _main_window_reduced


def _theme():
    values = dict(background_type='solid', background_color='#123456',
                  display_vertical_align=1, display_horizontal_align=2,
                  display_vertical_align_footer=0, display_horizontal_align_footer=0)
    for part, y, size in (('main', 80, 40), ('footer', 680, 12)):
        values.update({f'font_{part}_x': 20, f'font_{part}_y': y,
                       f'font_{part}_width': 984, f'font_{part}_height': 70,
                       f'font_{part}_name': 'Arial', f'font_{part}_size': size,
                       f'font_{part}_color': '#ffffff', f'font_{part}_bold': False,
                       f'font_{part}_italics': False})
    return SimpleNamespace(**values)


def test_pptx_has_editable_text_and_theme_geometry(tmp_path):
    from pptx import Presentation

    target = tmp_path / 'service.pptx'
    _write_pptx([PresentationPage(theme=_theme(), text='First line\nSecond line', footer='Credits')],
                target, 1024, 768)
    presentation = Presentation(str(target))
    assert len(presentation.slides) == 1
    assert [shape.text for shape in presentation.slides[0].shapes] == ['First line\nSecond line', 'Credits']
    assert presentation.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.name == 'Arial'


def test_odp_has_editable_text_and_master_page(tmp_path):
    target = tmp_path / 'service.odp'
    _write_odp([PresentationPage(theme=_theme(), text='Lyrics', footer='Credits')], target, 1024, 768)
    with ZipFile(target) as package:
        content = package.read('content.xml').decode('utf-8')
        styles = package.read('styles.xml').decode('utf-8')
        assert 'Lyrics' in content
        assert 'Credits' in content
        assert 'OpenLPMaster0' in content
        assert 'OpenLPMaster0' in styles


def test_service_pages_end_each_section_with_logo():
    theme = _theme()

    def item(title, lines):
        return SimpleNamespace(
            service_item_type=ServiceItemType.Text, title=title,
            get_theme_data=lambda: theme,
            rendered_slides=[{'text': '<strong>' + line + '</strong>',
                              'footer': '<b>Credits</b>'} for line in lines],
            display_slides=[{'text': line} for line in lines])

    pages = service_pages([item('First', ['Verse one', 'Verse two']),
                           item('Second', ['Another'])], insert_logo=True)
    assert [page.text for page in pages] == ['Verse one', 'Verse two', '', 'Another', '']
    assert [page.logo for page in pages] == [False, False, True, False, True]
    assert [page.section_name for page in pages] == ['First'] * 3 + ['Second'] * 2
    assert [page.section_index for page in pages] == [0, 0, 0, 1, 1]
    assert pages[0].footer == 'Credits'
    assert pages[0].html_text == '<strong>Verse one</strong>'


@pytest.mark.parametrize('thumbnail', ['icon', None, '', 'missing-thumbnail.png'])
def test_service_pages_without_thumbnail_use_title(thumbnail):
    if thumbnail == 'icon':
        pixmap = QtGui.QPixmap(16, 16)
        pixmap.fill(QtGui.QColor('#ff0000'))
        thumbnail = QtGui.QIcon(pixmap)
    theme = _theme()
    item = SimpleNamespace(service_item_type=ServiceItemType.Command, title='Welcome video',
                           get_theme_data=lambda: theme, slides=[{'image': thumbnail}])

    pages = service_pages([item], insert_logo=True)

    assert len(pages) == 2
    assert pages[0].image is None
    assert pages[0].theme is theme
    assert pages[0].text == 'Welcome video'
    assert pages[1].logo
    assert [page.section_index for page in pages] == [0, 0]
    assert [page.section_name for page in pages] == ['Welcome video', 'Welcome video']


@pytest.mark.parametrize('item_type, image_key', [
    (ServiceItemType.Image, 'path'), (ServiceItemType.Command, 'image')])
@pytest.mark.parametrize('path_type', [str, Path])
def test_service_pages_keep_thumbnail_files(tmp_path, item_type, image_key, path_type):
    thumbnail = tmp_path / 'thumbnail.png'
    image = QtGui.QImage(16, 16, QtGui.QImage.Format.Format_RGB32)
    image.fill(QtGui.QColor('#ff0000'))
    assert image.save(str(thumbnail))
    item = SimpleNamespace(service_item_type=item_type, title='Slide',
                           slides=[{image_key: path_type(thumbnail)}])

    pages = service_pages([item])

    assert len(pages) == 1
    assert pages[0].image == thumbnail


def test_command_without_source_uses_title_in_both_formats(tmp_path):
    from pptx import Presentation

    pixmap = QtGui.QPixmap(16, 16)
    pixmap.fill(QtGui.QColor('#ff0000'))
    item = SimpleNamespace(service_item_type=ServiceItemType.Command, title='Welcome video',
                           get_theme_data=_theme, slides=[{'image': QtGui.QIcon(pixmap)}])
    pages = service_pages([item])
    pptx = tmp_path / 'media.pptx'
    odp = tmp_path / 'media.odp'

    _write_pptx(pages, pptx, 1024, 768)
    _write_odp(pages, odp, 1024, 768)

    presentation = Presentation(str(pptx))
    assert len(presentation.slides) == 1
    assert presentation.slides[0].shapes[0].text == 'Welcome video'
    with ZipFile(odp) as package:
        assert b'Welcome video' in package.read('content.xml')


def _file_item(path, name, title, auto_start=False):
    return SimpleNamespace(name=name, title=title, service_item_type=ServiceItemType.Command,
                           slides=[{'image': QtGui.QIcon()}], get_frame_path=lambda: path,
                           is_capable=lambda capability: False, will_auto_start=auto_start)


@pytest.mark.parametrize('auto_start', [True, False])
def test_video_is_embedded_in_both_export_formats(tmp_path, auto_start):
    from pptx import Presentation
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    video = RESOURCE_PATH / 'media' / 'mp4_file.mp4'
    item = _file_item(video, 'media', 'Welcome video', auto_start)
    pages = service_pages([item], native_pptx=True)
    target = tmp_path / 'video.pptx'
    _write_pptx(pages, target, 1024, 768)
    slide = Presentation(target).slides[0]
    media = slide.part.part_related_by(RT.MEDIA)
    assert media.blob == video.read_bytes()
    assert media.content_type == 'video/mp4'
    assert slide.part.part_related_by(RT.VIDEO) is media
    conditions = slide._element.xpath('.//p:video/p:cMediaNode/p:cTn/p:stCondLst/p:cond')
    assert conditions[0].get('delay') == ('0' if auto_start else 'indefinite')

    target = tmp_path / 'video.odp'
    _write_odp(pages, target, 1024, 768)
    with ZipFile(target) as archive:
        assert archive.read('Media/0.mp4') == video.read_bytes()
        assert b'application/vnd.sun.star.media' in archive.read('content.xml')
        assert b'Media/0.mp4' in archive.read('content.xml')
        assert b'video/mp4' in archive.read('META-INF/manifest.xml')


def _source_pptx(tmp_path):
    from lxml import etree
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches, Pt

    source = Presentation()
    source.slide_width, source.slide_height = Inches(10), Inches(5.625)
    first = source.slides.add_slide(source.slide_layouts[6])
    second = source.slides.add_slide(source.slide_layouts[6])
    text = first.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    run = text.text_frame.paragraphs[0].add_run()
    run.text, run.font.bold, run.font.size = 'Editable original text', True, Pt(32)
    text.click_action.target_slide = second
    second.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text = 'Second original slide'
    first.notes_slide.notes_text_frame.text = 'Original speaker notes'
    image_path = tmp_path / 'original.png'
    image = QtGui.QImage(640, 360, QtGui.QImage.Format.Format_RGB32)
    image.fill(QtGui.QColor('#ff1234'))
    image.save(str(image_path))
    first.shapes.add_picture(str(image_path), Inches(1), Inches(2), Inches(2), Inches(1))
    chart_data = CategoryChartData()
    chart_data.categories = ['One', 'Two']
    chart_data.add_series('Original data', (1, 2))
    first.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5), Inches(1), Inches(4), Inches(3), chart_data)
    second.shapes.add_movie(str(RESOURCE_PATH / 'media' / 'mp4_file.mp4'), 0, Inches(2), Inches(4), Inches(2),
                            mime_type='video/mp4')
    transition = etree.SubElement(first._element,
                                  '{http://schemas.openxmlformats.org/presentationml/2006/main}transition', spd='slow')
    etree.SubElement(transition, '{http://schemas.openxmlformats.org/presentationml/2006/main}fade')
    path = tmp_path / 'original.pptx'
    source.save(path)
    return path


def test_native_pptx_slides_keep_editable_content_assets_and_internal_links(tmp_path, monkeypatch):
    import posixpath
    from lxml import etree
    from pptx import Presentation
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    path = _source_pptx(tmp_path)
    original = Presentation(path)
    item = _file_item(path, 'presentations', 'Sermon')
    monkeypatch.setattr('openlp.core.lib.presentationexport._background',
                        lambda *args: ('#000000', None))
    # A stale cache lists only one thumbnail, but both original slides must be exported.
    pages = service_pages([item, item], insert_logo=True, native_pptx=True)
    assert len(pages) == 6
    target = tmp_path / 'service.pptx'
    _write_pptx(pages, target, 1024, 768)
    exported = Presentation(target)
    assert (exported.slide_width, exported.slide_height) == (original.slide_width, original.slide_height)
    assert len(exported.slides) == 6
    for first_index in (0, 3):
        first, second = exported.slides[first_index], exported.slides[first_index + 1]
        assert first.shapes[0].text == 'Editable original text'
        assert first.shapes[0].text_frame.paragraphs[0].runs[0].font.bold
        assert first.shapes[0].click_action.target_slide == second
        assert first.shapes[1].image.blob == original.slides[0].shapes[1].image.blob
        assert first.shapes[2].chart.series[0].values == (1.0, 2.0)
        assert first.notes_slide.notes_text_frame.text == 'Original speaker notes'
        assert first._element.xpath('./p:transition/p:fade')
        assert second.part.part_related_by(RT.MEDIA).blob == (RESOURCE_PATH / 'media' / 'mp4_file.mp4').read_bytes()
        original_master = etree.fromstring(original.slides[0].slide_layout.slide_master.part.blob)
        copied_master = etree.fromstring(first.slide_layout.slide_master.part.blob)
        for master in (original_master, copied_master):
            for layout in master.xpath('//*[local-name()="sldLayoutId"]'):
                layout.attrib.pop('id')
        assert etree.tostring(copied_master) == etree.tostring(original_master)
    assert exported.slides[0].part is not exported.slides[3].part
    with ZipFile(target) as archive:
        names = archive.namelist()
        assert len(names) == len(set(names))
        for name in names:
            if not name.endswith('.rels'):
                continue
            for rel in etree.fromstring(archive.read(name)):
                if rel.get('TargetMode') != 'External':
                    base = posixpath.dirname(posixpath.dirname(name))
                    dependency = posixpath.normpath(posixpath.join(base, rel.get('Target'))).lstrip('/')
                    assert dependency in names, (name, dependency)
        sections = etree.fromstring(archive.read('ppt/presentation.xml')).xpath(
            '//*[local-name()="section"]')
        assert [section.get('name') for section in sections] == ['Sermon', 'Sermon']
        assert [len(section[0]) for section in sections] == [3, 3]
        master_ids = etree.fromstring(archive.read('ppt/presentation.xml')).xpath(
            '//*[local-name()="sldMasterId"]/@id')
        for name in names:
            if 'slideMaster' in name and name.endswith('.xml'):
                master_ids.extend(etree.fromstring(archive.read(name)).xpath('//*[local-name()="sldLayoutId"]/@id'))
        assert len(master_ids) == len(set(master_ids))


def test_missing_media_fails_instead_of_silently_exporting_a_title(tmp_path):
    item = _file_item(tmp_path / 'missing.mp4', 'media', 'Missing video')
    with pytest.raises(FileNotFoundError, match='media file is missing'):
        service_pages([item], native_pptx=True)


def test_audio_media_is_packaged_as_audio(tmp_path):
    from pptx import Presentation
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    source = RESOURCE_PATH / 'media' / 'mp3_file.mp3'
    pages = service_pages([_file_item(source, 'media', 'Music', auto_start=True)], native_pptx=True)
    target = tmp_path / 'audio.pptx'
    _write_pptx(pages, target, 1024, 768)
    slide = Presentation(target).slides[0]
    assert slide.part.part_related_by(RT.AUDIO).blob == source.read_bytes()
    assert slide._element.xpath('.//a:audioFile')
    assert slide._element.xpath('.//p:audio/p:cMediaNode/p:cTn/p:stCondLst/p:cond')[0].get('delay') == '0'


def test_video_poster_uses_decoded_frame_and_service_start_time(tmp_path, monkeypatch):
    from pptx import Presentation
    from openlp.core.lib import mediaposter

    decoded = QtGui.QImage(1920, 1080, QtGui.QImage.Format.Format_RGB32)
    decoded.fill(QtGui.QColor('#247b39'))
    calls = []

    def capture(path, position_ms):
        calls.append((path, position_ms))
        return decoded

    monkeypatch.setattr(mediaposter, 'video_frame', capture)
    media_path = RESOURCE_PATH / 'media' / 'mp4_file.mp4'
    item = _file_item(media_path, 'media', 'Video')
    item.start_time = 5000
    pages = service_pages([item, item], native_pptx=True)
    target = tmp_path / 'poster.pptx'
    _write_pptx(pages, target, 1920, 1080)

    assert calls == [(media_path, 5000)]  # Repeated media reuse the poster within this export.
    for slide in Presentation(target).slides:
        image = QtGui.QImage.fromData(slide.shapes[0].poster_frame.blob)
        assert (image.width(), image.height()) == (1920, 1080)
        assert image.pixelColor(960, 540).name() == '#247b39'
        assert image.pixelColor(200, 200).name() == '#247b39'


def test_video_poster_preserves_aspect_ratio_and_has_high_resolution_fallback(monkeypatch):
    from openlp.core.lib import mediaposter, presentationexport

    page = PresentationPage(media=Path('video.mp4'))
    portrait = QtGui.QImage(600, 1200, QtGui.QImage.Format.Format_RGB32)
    portrait.fill(QtGui.QColor('#abcdef'))
    monkeypatch.setattr(mediaposter, 'video_frame', lambda *args, **kwargs: portrait)
    poster = QtGui.QImage.fromData(presentationexport._media_poster(page, 16 / 9))
    assert poster.pixelColor(960, 540).name() == '#abcdef'
    assert poster.pixelColor(200, 540).name() == '#000000'

    monkeypatch.setattr(mediaposter, 'video_frame', lambda *args, **kwargs: QtGui.QImage())
    fallback = QtGui.QImage.fromData(presentationexport._media_poster(page, 16 / 9))
    assert (fallback.width(), fallback.height()) == (1920, 1080)
    assert fallback.pixelColor(960, 540).name() == '#dddddd'
    assert fallback.pixelColor(960, 300).name() == '#000000'


def test_export_mixed_service_keeps_native_slides_video_and_logo_order(tmp_path, monkeypatch):
    from PySide6 import QtCore
    from pptx import Presentation
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from openlp.core.lib import presentationexport

    song = SimpleNamespace(service_item_type=ServiceItemType.Text, title='Song', get_theme_data=_theme,
                           rendered_slides=[{'text': 'Editable lyrics'}], display_slides=[{'text': 'Editable lyrics'}])
    media = _file_item(RESOURCE_PATH / 'media' / 'mp4_file.mp4', 'media', 'Video')
    source = _file_item(_source_pptx(tmp_path), 'presentations', 'Sermon')
    screen = SimpleNamespace(display_geometry=QtCore.QRect(0, 0, 1920, 1080))
    monkeypatch.setattr(presentationexport, 'ScreenList',
                        lambda: SimpleNamespace(current=screen))
    monkeypatch.setattr(presentationexport, '_background', lambda *args: ('#000000', None))
    target = tmp_path / 'mixed.pptx'
    presentationexport.export_presentation([song, media, source], target, insert_logo=True)
    slides = Presentation(target).slides
    assert len(slides) == 7
    assert slides[0].shapes[0].text == 'Editable lyrics'
    assert slides[2].part.part_related_by(RT.MEDIA).blob == media.get_frame_path().read_bytes()
    assert slides[4].shapes[0].text == 'Editable original text'
    assert slides[5].shapes[0].text == 'Second original slide'
    assert all(len(slides[index].shapes) == 0 for index in (1, 3, 6))


def test_different_source_slide_sizes_are_fitted_without_rasterising(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches, Pt

    paths = []
    for index, size in enumerate(((10, 5), (5, 5))):
        source = Presentation()
        source.slide_width, source.slide_height = (Inches(value) for value in size)
        slide = source.slides.add_slide(source.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
        box.text = 'Native text'
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
        path = tmp_path / f'source{index}.pptx'
        source.save(path)
        paths.append(path)
    pages = service_pages([_file_item(path, 'presentations', path.stem) for path in paths], native_pptx=True)
    target = tmp_path / 'mixed-sizes.pptx'
    _write_pptx(pages, target, 1920, 1080)
    result = Presentation(target)
    box = result.slides[1].shapes[0]
    assert box.text == 'Native text'
    assert box.left == Inches(3.5)
    assert box.top == Inches(1)
    assert box.width == Inches(2)
    assert box.text_frame.paragraphs[0].runs[0].font.size == Pt(24)


def test_service_sections_are_exported_to_pptx_and_odp(tmp_path, monkeypatch):
    from lxml import etree

    theme = _theme()

    def item(title):
        return SimpleNamespace(service_item_type=ServiceItemType.Text, title=title,
                               get_theme_data=lambda: theme,
                               rendered_slides=[{'text': title, 'footer': ''}],
                               display_slides=[{'text': title}])

    monkeypatch.setattr('openlp.core.lib.presentationexport._background',
                        lambda theme, width, height, logo, logo_settings=None: ('#000000', None))
    pages = service_pages([item('Pre-Service'), item('How Great Thou Art'),
                           item('How Great Thou Art')], insert_logo=True)
    pptx = tmp_path / 'sections.pptx'
    odp = tmp_path / 'sections.odp'
    _write_pptx(pages, pptx, 1024, 768)
    _write_odp(pages, odp, 1024, 768)

    pptx_ns = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
               'p14': 'http://schemas.microsoft.com/office/powerpoint/2010/main'}
    with ZipFile(pptx) as package:
        root = etree.fromstring(package.read('ppt/presentation.xml'))
    slide_ids = root.xpath('./p:sldIdLst/p:sldId/@id', namespaces=pptx_ns)
    sections = root.xpath('./p:extLst/p:ext/p14:sectionLst/p14:section', namespaces=pptx_ns)
    assert [section.get('name') for section in sections] == [
        'Pre-Service', 'How Great Thou Art', 'How Great Thou Art']
    assert [[slide_id.get('id') for slide_id in section.xpath(
        './p14:sldIdLst/p14:sldId', namespaces=pptx_ns)]
            for section in sections] == [slide_ids[0:2], slide_ids[2:4], slide_ids[4:6]]
    assert len({section.get('id') for section in sections}) == 3

    odp_ns = {'office': 'urn:oasis:names:tc:opendocument:xmlns:office:1.0',
              'draw': 'urn:oasis:names:tc:opendocument:xmlns:drawing:1.0',
              'loext': 'urn:org:documentfoundation:names:experimental:office:xmlns:loext:1.0'}
    with ZipFile(odp) as package:
        root = etree.fromstring(package.read('content.xml'))
    slide_names = root.xpath('./office:body/office:presentation/draw:page/@draw:name', namespaces=odp_ns)
    sections = root.xpath('./office:body/office:presentation/loext:section-list/loext:section',
                          namespaces=odp_ns)
    assert [section.get(f'{{{odp_ns["loext"]}}}name') for section in sections] == [
        'Pre-Service', 'How Great Thou Art', 'How Great Thou Art']
    assert [[slide.get(f'{{{odp_ns["loext"]}}}name') for slide in section]
            for section in sections] == [slide_names[0:2], slide_names[2:4], slide_names[4:6]]


def test_logo_slide_is_packaged_in_both_formats(tmp_path, monkeypatch):
    from pptx import Presentation

    logo = tmp_path / 'logo.png'
    image = QtGui.QImage(16, 16, QtGui.QImage.Format.Format_RGB32)
    image.fill(QtGui.QColor('#ff0000'))
    assert image.save(str(logo))
    settings = SimpleNamespace(value=lambda key: {
        'core/logo background color': '#000000',
        'core/logo file': logo,
        'core/logo fit': 'contain'}[key])
    monkeypatch.setattr('openlp.core.lib.presentationexport.Registry',
                        lambda: SimpleNamespace(get=lambda key: settings))
    pages = [PresentationPage(theme=_theme(), text='First'), PresentationPage(logo=True),
             PresentationPage(theme=_theme(), text='Second')]
    pptx = tmp_path / 'with-logo.pptx'
    odp = tmp_path / 'with-logo.odp'
    _write_pptx(pages, pptx, 1024, 768)
    _write_odp(pages, odp, 1024, 768)
    assert len(Presentation(str(pptx)).slides) == 3
    with ZipFile(odp) as package:
        assert package.read('content.xml').count(b'draw:page ') == 3
        assert any(name.startswith('Pictures/') for name in package.namelist())


def test_formatting_tags_become_editable_runs_in_both_formats(tmp_path):
    from pptx import Presentation

    markup = ('<strong>Bold</strong> '
              '<span style="-webkit-text-fill-color:red">Red</span><br>'
              '<span style="font-size:150%; font-family:Georgia; text-decoration:underline">Large</span>'
              '<sup>2</sup>')
    page = PresentationPage(theme=_theme(), text='Bold Red\nLarge2', html_text=markup,
                            footer='Credits', html_footer='<span style="color:blue">Credits</span>')
    lines = _styled_lines(markup, _base_text_style(page.theme, 'main'))
    assert [[value for value, _ in line] for line in lines] == [['Bold', ' ', 'Red'], ['Large', '2']]
    pptx = tmp_path / 'formatted.pptx'
    odp = tmp_path / 'formatted.odp'
    _write_pptx([page], pptx, 1024, 768)
    _write_odp([page], odp, 1024, 768)
    shapes = Presentation(str(pptx)).slides[0].shapes
    paragraphs = shapes[0].text_frame.paragraphs
    assert paragraphs[0].runs[0].font.bold
    assert paragraphs[0].runs[2].font.color.rgb == (255, 0, 0)
    assert paragraphs[1].runs[0].font.name == 'Georgia'
    assert paragraphs[1].runs[0].font.underline
    assert paragraphs[1].runs[0].font.size > paragraphs[0].runs[0].font.size
    assert paragraphs[1].runs[1]._r.get_or_add_rPr().get('baseline') == '30000'
    assert shapes[1].text_frame.paragraphs[0].runs[0].font.color.rgb == (0, 0, 255)
    with ZipFile(odp) as package:
        content = package.read('content.xml').decode('utf-8')
        assert 'font-family="Georgia"' in content
        assert 'text-underline-style="solid"' in content
        assert 'color="#ff0000"' in content
        assert 'text-position="super 58%"' in content


def test_openlp_and_custom_formatting_tags_are_expanded(monkeypatch):
    tags = [
        {'start tag': '{st}', 'end tag': '{/st}', 'start html': '<strong>', 'end html': '</strong>'},
        {'start tag': '{r}', 'end tag': '{/r}',
         'start html': '<span style="-webkit-text-fill-color:red">', 'end html': '</span>'},
        {'start tag': '{large}', 'end tag': '{/large}',
         'start html': '<span style="font-size:150%; font-family:Georgia">', 'end html': '</span>'},
    ]
    monkeypatch.setattr(FormattingTags, 'get_html_tags', lambda: tags)
    monkeypatch.setattr(render_module, 'remove_chords', lambda value: value)
    markup = render_module.render_tags('{st}Bold{/st} {r}Red{/r} {large}Large{/large}')
    runs = _styled_lines(markup, _base_text_style(_theme(), 'main'))[0]
    assert runs[0][1].bold
    assert runs[2][1].color == '#ff0000'
    assert runs[4][1].family == 'Georgia'
    assert runs[4][1].size == 60


def test_presentation_metadata_identifies_openlp(tmp_path):
    from pptx import Presentation

    page = PresentationPage(theme=_theme(), text='Lyrics')
    pptx = tmp_path / 'Sunday.pptx'
    odp = tmp_path / 'Sunday.odp'
    _write_pptx([page], pptx, 1024, 768)
    _write_odp([page], odp, 1024, 768)
    properties = Presentation(str(pptx)).core_properties
    assert properties.title == 'Sunday'
    assert properties.comments == 'Exported from OpenLP+'
    assert properties.author == 'OpenLP+'
    assert properties.last_modified_by == 'OpenLP+'
    with ZipFile(pptx) as package:
        assert b'generated using python-pptx' not in package.read('docProps/core.xml')
    with ZipFile(odp) as package:
        metadata = package.read('meta.xml')
        assert b'Exported from OpenLP+' in metadata


def test_presentation_filename_matches_saved_service():
    manager = SimpleNamespace(file_name=lambda: Path('/services/Sunday Evening.oszl'))
    assert ServiceManager.suggested_presentation_file_path(manager, '.pptx') == \
        Path('/services/Sunday Evening.pptx')


def test_presentation_filename_uses_service_name_pattern(monkeypatch, tmp_path):
    class FixedDatetime:
        @staticmethod
        def now():
            return datetime(2026, 9, 23, 10, 0)

    monkeypatch.setattr('openlp.core.ui.servicemanager.datetime', FixedDatetime)
    values = {'advanced/default service enabled': True,
              'advanced/default service day': 7,
              'advanced/default service name': 'Service %Y-%m-%d %H-%M',
              'servicemanager/last directory': tmp_path}
    manager = SimpleNamespace(file_name=lambda: None,
                              settings=SimpleNamespace(value=lambda key: values[key]))
    manager._default_service_file_path = lambda: ServiceManager._default_service_file_path(manager)
    assert ServiceManager.suggested_presentation_file_path(manager, '.odp') == \
        tmp_path / 'Service 2026-09-23 10-00.odp'


def test_presentation_filename_without_default_name_uses_service_label(tmp_path):
    values = {'advanced/default service enabled': False,
              'servicemanager/last directory': tmp_path}
    manager = SimpleNamespace(file_name=lambda: None,
                              settings=SimpleNamespace(value=lambda key: values[key]))
    assert ServiceManager.suggested_presentation_file_path(manager, '.pptx') == tmp_path / 'Service.pptx'


def test_export_dialog_uses_requested_format_and_suggested_path(monkeypatch, tmp_path):
    from openlp.core.lib import presentationexport

    saved = []
    dialogs = []
    suggested = tmp_path / 'Service 2026-09-23.odp'

    def choose_file(*args):
        dialogs.append(args)
        return suggested, ''

    monkeypatch.setattr(FileDialog, 'getSaveFileName', choose_file)
    monkeypatch.setattr(presentationexport, 'export_presentation',
                        lambda items, path, insert_logo, **kwargs: saved.append((path, insert_logo)))
    presentationexport.export_items_dialog(None, [SimpleNamespace(title='Song')], '.odp', suggested)
    assert saved == [(suggested, False)]
    assert dialogs[0][2] == suggested
    assert dialogs[0][3] == 'OpenDocument Presentation (*.odp)'


@pytest.mark.parametrize('fails', [False, True])
def test_export_progress_dialog_is_modal_and_closes_before_errors(tmp_path, monkeypatch, fails):
    from PySide6 import QtCore, QtWidgets
    from openlp.core.lib import presentationexport
    from openlp.core.ui import presentationexportdialog

    dialogs = []
    errors = []
    original_dialog = presentationexportdialog.ExportProgressDialog

    def create_dialog(parent):
        dialog = original_dialog(parent)
        dialogs.append(dialog)
        return dialog

    def export(items, path, insert_logo, *, progress):
        dialog = dialogs[0]
        assert dialog.isVisible()
        assert dialog.windowModality() == QtCore.Qt.WindowModality.ApplicationModal
        progress('Embedding media\nWelcome video', 1, 3)
        assert dialog.value() == 1
        assert dialog.maximum() == 3
        assert dialog.labelText() == 'Embedding media\nWelcome video'
        dialog.reject()
        assert dialog.isVisible()
        progress('Writing file', 0, 0)
        assert dialog.minimum() == dialog.maximum() == 0
        if fails:
            raise OSError('Disk full')

    def show_error(parent, title, message):
        assert not dialogs[0].isVisible()
        errors.append(message)

    monkeypatch.setattr(presentationexportdialog, 'ExportProgressDialog', create_dialog)
    monkeypatch.setattr(FileDialog, 'getSaveFileName', lambda *args: (tmp_path / 'service.pptx', ''))
    monkeypatch.setattr(presentationexport, 'export_presentation', export)
    monkeypatch.setattr(QtWidgets.QMessageBox, 'critical', show_error)
    presentationexport.export_items_dialog(None, [SimpleNamespace(title='Song')], '.pptx', tmp_path)
    assert not dialogs[0].isVisible()
    assert errors == (['Disk full'] if fails else [])


def test_cancelled_save_dialog_does_not_start_progress(tmp_path, monkeypatch):
    from openlp.core.lib import presentationexport
    from openlp.core.ui import presentationexportdialog

    def unexpected(*args):
        pytest.fail('Cancelling the file chooser must not start an export')

    monkeypatch.setattr(FileDialog, 'getSaveFileName', lambda *args: (None, ''))
    monkeypatch.setattr(presentationexportdialog, 'ExportProgressDialog', unexpected)
    presentationexport.export_items_dialog(None, [SimpleNamespace(title='Song')], '.pptx', tmp_path)


def test_export_worker_keeps_gui_responsive_and_returns_errors():
    import threading
    from PySide6 import QtCore
    from openlp.core.ui.presentationexportdialog import write_with_progress

    gui_thread = threading.get_ident()
    tick = threading.Event()
    updates = []

    def action(progress):
        assert threading.get_ident() != gui_thread
        progress('Writing file', 0, 0)
        assert tick.wait(2), 'The GUI timer must run while the writer is busy'
        raise OSError('Write failed')

    timer = QtCore.QTimer()
    timer.timeout.connect(tick.set)
    timer.start(10)
    try:
        with pytest.raises(OSError, match='Write failed'):
            write_with_progress(action, lambda *args: updates.append((threading.get_ident(), args)))
    finally:
        timer.stop()
    assert updates == [(gui_thread, ('Writing file', 0, 0))]


@pytest.mark.parametrize('suffix', ['.pptx', '.odp'])
def test_export_with_progress_writes_real_output_on_worker(tmp_path, monkeypatch, registry, suffix):
    import threading
    from PySide6 import QtCore
    from openlp.core.lib import presentationexport

    gui_thread = threading.get_ident()
    updates = []
    registry.register('settings', SimpleNamespace(value=lambda key: {
        'core/logo background color': '#000000', 'core/logo file': None, 'core/logo fit': 'contain'}[key]))
    screen = SimpleNamespace(display_geometry=QtCore.QRect(0, 0, 1920, 1080))
    monkeypatch.setattr(presentationexport, 'ScreenList', lambda: SimpleNamespace(current=screen))
    media = _file_item(RESOURCE_PATH / 'media' / 'mp4_file.mp4', 'media', 'Video')
    items = [media]
    if suffix == '.pptx':
        items.append(_file_item(_source_pptx(tmp_path), 'presentations', 'Sermon'))
    target = tmp_path / ('service' + suffix)

    def progress(*args):
        assert threading.get_ident() == gui_thread
        updates.append(args)

    presentationexport.export_presentation(items, target, insert_logo=True, progress=progress)

    assert target.is_file()
    assert updates[0] == ('Preparing slides\nVideo', 0, len(items))
    assert any(message.startswith('Embedding media') for message, _, _ in updates)
    assert updates[-1] == ('Writing file\n' + target.name, 0, 0)
    with ZipFile(target) as archive:
        assert archive.testzip() is None
    if suffix == '.pptx':
        from pptx import Presentation

        assert any(message.startswith('Copying original slides') for message, _, _ in updates)
        assert len(Presentation(target).slides) == 5


def test_export_menu_has_separate_presentation_formats(request):
    main_window = request.getfixturevalue('presentation_main_window')
    menu = main_window.export_presentation_menu
    assert menu.title() == '&Presentation'
    assert [action.text() for action in menu.actions()] == ['PowerPoint (PPTX)', 'OpenDocument (ODP)']
    assert menu.menuAction() in main_window.file_export_menu.actions()
    main_window.export_powerpoint_item.trigger()
    main_window.export_opendocument_item.trigger()
    calls = main_window.service_manager_contents.on_export_presentation.call_args_list
    assert [call.args[0] for call in calls] == ['.pptx', '.odp']


def test_pptx_keeps_editable_design_layers_effects_and_embedded_images(tmp_path):
    import base64
    from pptx import Presentation
    from PySide6.QtCore import QSize
    from openlp.core.lib.presentationexport import _image_bytes
    from openlp.core.lib.slidedesign import apply_design

    picture = QtGui.QImage(80, 40, QtGui.QImage.Format.Format_ARGB32)
    picture.fill(QtGui.QColor('#ff0000'))
    uri = 'data:image/png;base64,' + base64.b64encode(_image_bytes(picture)).decode()
    design = {'version': 1, 'canvas_width': 1920, 'canvas_height': 1080,
              'lyrics': {'x': 10, 'y': 10, 'width': 80, 'height': 20, 'outline': True,
                         'outline_width': 2, 'shadow': True, 'rotation': 5},
              'background': {'type': 'gradient', 'colour': '#123456', 'end_colour': '#654321',
                             'image': uri, 'image_opacity': 30},
              'order': ['ellipse', 'lyrics', 'caption', 'image'],
              'elements': [
                  {'id': 'ellipse', 'kind': 'ellipse', 'x': 10, 'y': 30, 'width': 20, 'height': 20,
                   'fill': '#00ff00', 'opacity': 40, 'outline': True, 'outline_width': 4, 'shadow': True},
                  {'id': 'caption', 'kind': 'text', 'text': 'Editable caption', 'x': 35, 'y': 60,
                   'width': 30, 'height': 20, 'font_size': 36, 'opacity': 60,
                   'outline': True, 'shadow': True, 'shadow_x': -8, 'flip_vertical': True},
                  {'id': 'image', 'kind': 'image', 'src': uri, 'x': 70, 'y': 50, 'width': 20, 'height': 30,
                   'opacity': 50, 'flip_horizontal': True, 'outline': True, 'outline_width': 3, 'shadow': True}]}
    theme = apply_design(_theme(), design, QSize(1920, 1080))
    destination = tmp_path / 'rich.pptx'
    _write_pptx([PresentationPage(theme=theme, text='Slide lyrics', footer='')], destination, 1920, 1080)
    deck = Presentation(str(destination))
    shapes = deck.slides[0].shapes
    assert [shape.name for shape in shapes][1:] == ['Ellipse', 'Slide text', 'Text box', 'Image']
    assert shapes[2].text == 'Slide lyrics'
    assert shapes[3].text == 'Editable caption'
    assert shapes[2].rotation == 5
    assert shapes[1].width == pytest.approx(deck.slide_width * .2, abs=1)
    assert shapes[3]._element.xpath('./p:spPr/a:scene3d/a:camera/a:rot')[0].get('lat') == '10800000'
    assert shapes[4]._element.xpath('./p:spPr/a:xfrm')[0].get('flipH') == '1'
    assert shapes[4]._element.xpath('./p:blipFill/a:blip/a:alphaModFix')[0].get('amt') == '50000'
    assert shapes[1]._element.xpath('./p:spPr/a:solidFill/a:srgbClr/a:alpha')[0].get('val') == '40000'
    assert shapes[1]._element.xpath('./p:spPr/a:effectLst/a:outerShdw')
    assert shapes[3]._element.xpath('.//a:rPr/a:ln/a:solidFill')
    assert shapes[3]._element.xpath('.//a:rPr/a:effectLst/a:outerShdw')
    assert len(shapes) == 5  # Background plus four editable objects, without a footer.
